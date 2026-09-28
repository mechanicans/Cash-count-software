from datetime import datetime, timedelta
from tkinter import filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from database import load_records_between


EXPORT_OPTIONS = {
    "3 Days": ("days", 3),
    "7 Days": ("days", 7),
    "15 Days": ("days", 15),
    "30 Days": ("days", 30),
    "3 Months": ("months", 3),
    "6 Months": ("months", 6),
    "7 Months": ("months", 7),
}


def subtract_months(source_date, months):

    year = source_date.year
    month = source_date.month - months

    while month <= 0:
        month += 12
        year -= 1

    # First day of resulting month is not what we want.
    # We preserve the day where possible.
    import calendar

    last_day = calendar.monthrange(year, month)[1]
    day = min(source_date.day, last_day)

    return source_date.replace(
        year=year,
        month=month,
        day=day
    )


def calculate_start_date(end_date, range_name):

    range_type, value = EXPORT_OPTIONS[range_name]

    if range_type == "days":
        # Inclusive date range:
        # 3 Days means end date + previous 2 days.
        return end_date - timedelta(days=value - 1)

    if range_type == "months":
        start_date = subtract_months(end_date, value)

        # Inclusive range convention.
        return start_date + timedelta(days=1)

    raise ValueError("Invalid export range.")


def export_cash_ledger(end_date_string, range_name, parent=None):

    try:
        end_date = datetime.strptime(
            end_date_string,
            "%Y-%m-%d"
        ).date()

        start_date = calculate_start_date(
            end_date,
            range_name
        )

        rows = load_records_between(
            start_date.strftime("%Y-%m-%d"),
            end_date.strftime("%Y-%m-%d")
        )

        if not rows:
            messagebox.showinfo(
                "Export Cash Ledger",
                "No saved records were found in the selected period.",
                parent=parent
            )
            return False

        file_path = filedialog.asksaveasfilename(
            parent=parent,
            title="Save Cash Ledger Excel File",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Workbook", "*.xlsx")
            ],
            initialfile=(
                f"Cash_Ledger_"
                f"{start_date.strftime('%Y-%m-%d')}_to_"
                f"{end_date.strftime('%Y-%m-%d')}.xlsx"
            )
        )

        if not file_path:
            return False

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Cash Ledger"

        headers = [
            "Date",
            "₹2000 Count",
            "₹500 Count",
            "₹200 Count",
            "₹100 Count",
            "₹50 Count",
            "₹20 Count",
            "₹10 Count",
            "₹5 Count",
            "₹2 Count",
            "₹1 Count",
            "Expense",
            "Remarks",
            "Total Cash",
            "Net Cash",
        ]

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        for row in rows:
            sheet.append(row)

        # Format date column
        for cell in sheet["A"][1:]:
            cell.alignment = Alignment(horizontal="center")

        # Format numeric columns
        for row in sheet.iter_rows(
            min_row=2,
            min_col=2,
            max_col=15
        ):
            for cell in row:
                cell.alignment = Alignment(horizontal="center")

        # Freeze header
        sheet.freeze_panes = "A2"

        # Auto filter
        sheet.auto_filter.ref = sheet.dimensions

        # Adjust column widths
        for column_cells in sheet.columns:

            maximum_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                maximum_length = max(
                    maximum_length,
                    len(value)
                )

            sheet.column_dimensions[column_letter].width = min(
                maximum_length + 3,
                25
            )

        workbook.save(file_path)

        messagebox.showinfo(
            "Export Successful",
            (
                f"{len(rows)} saved record(s) exported successfully.\n\n"
                f"Period:\n"
                f"{start_date.strftime('%d-%m-%Y')} to "
                f"{end_date.strftime('%d-%m-%Y')}"
            ),
            parent=parent
        )

        return True

    except Exception as error:

        messagebox.showerror(
            "Export Error",
            f"Could not export the cash ledger.\n\n{error}",
            parent=parent
        )

        return False